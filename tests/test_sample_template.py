"""Format assertions for the generated synthetic .docx template.

These prove the generator reproduces the reference's *format* (US-Letter, RTL,
teal/grey/white fills, exact grids, blue 18pt headings, page-number footer) and —
critically for the privacy rule — that the data rows hold docxtpl Jinja tags, not
any real names. The committed artifact is the script; the .docx is generated here.
"""

from __future__ import annotations

import struct
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from meeting_minutes import sample_template as st


@pytest.fixture(scope="session")
def template_path(tmp_path_factory) -> str:
    out = tmp_path_factory.mktemp("tmpl") / "sample.docx"
    return str(st.build(out))


@pytest.fixture(scope="session")
def doc(template_path):
    return Document(template_path)


def _fill(cell) -> str | None:
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def _grid(table) -> list[int]:
    return [int(gc.get(qn("w:w"))) for gc in table._tbl.tblGrid.findall(qn("w:gridCol"))]


def _all_text(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )


# --- page / section ---------------------------------------------------------

def test_page_is_us_letter_with_half_inch_margins(doc) -> None:
    s = doc.sections[0]
    assert (s.page_width, s.page_height) == (Twips(st.PAGE_W_TW), Twips(st.PAGE_H_TW))
    assert s.top_margin == s.bottom_margin == s.left_margin == s.right_margin == Twips(st.MARGIN_TW)
    assert s.header_distance == s.footer_distance == Twips(st.HEADER_FOOTER_TW)


def test_section_is_rtl(doc) -> None:
    assert doc.sections[0]._sectPr.find(qn("w:bidi")) is not None


# --- headings ---------------------------------------------------------------

def test_heading_style_is_blue_18pt_bold(doc) -> None:
    style = doc.styles["Heading 1"]
    assert style.font.color.rgb == RGBColor(0x00, 0x70, 0xB9)
    assert style.font.size == Pt(18)
    assert style.font.bold is True


def test_three_section_headings_present(doc) -> None:
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["قائمة الحضور", "نقاط نقاش الاجتماع", "نتائج الاجتماع"]


# --- tables: structure, grids, RTL -----------------------------------------

def test_four_tables_in_order(doc) -> None:
    assert len(doc.tables) == 4


def test_all_tables_are_rtl_bidi_visual(doc) -> None:
    for t in doc.tables:
        assert t._tbl.tblPr.find(qn("w:bidiVisual")) is not None


def test_table_grid_widths_match_reference(doc) -> None:
    _title, attendees, summary, outcomes = doc.tables
    assert _grid(attendees) == st.GRID_ATTEND == [447, 5386, 4957]
    assert _grid(summary) == st.GRID_SUMMARY == [10790]
    assert _grid(outcomes) == st.GRID_OUTCOME == [446, 5954, 2551, 1839]


# --- fills ------------------------------------------------------------------

def test_attendees_header_teal_index_grey_data_white(doc) -> None:
    attendees = doc.tables[1]
    assert all(_fill(c) == st.TEAL for c in attendees.rows[0].cells)
    data = attendees.rows[2].cells  # rows[1] is the {%tr for%} marker row
    assert _fill(data[0]) == st.GREY  # # column
    assert _fill(data[1]) == st.WHITE and _fill(data[2]) == st.WHITE


def test_summary_header_teal_body_white(doc) -> None:
    summary = doc.tables[2]
    assert _fill(summary.rows[0].cells[0]) == st.TEAL
    assert _fill(summary.rows[1].cells[0]) == st.WHITE


def test_outcomes_header_teal_index_grey_data_white(doc) -> None:
    outcomes = doc.tables[3]
    assert all(_fill(c) == st.TEAL for c in outcomes.rows[0].cells)
    data = outcomes.rows[2].cells  # rows[1] is the {%tr for%} marker row
    assert _fill(data[0]) == st.GREY
    assert all(_fill(c) == st.WHITE for c in data[1:])


# --- footer / body ----------------------------------------------------------

def test_footer_has_page_field(doc) -> None:
    footer = doc.sections[0].footer
    instrs = [f.get(qn("w:instr")) for p in footer.paragraphs for f in p._p.findall(qn("w:fldSimple"))]
    assert "PAGE" in instrs


def test_body_has_thanks_line(doc) -> None:
    assert any(p.text.strip() == st.THANKS for p in doc.paragraphs)


def test_header_embeds_an_image(template_path) -> None:
    media = [n for n in zipfile.ZipFile(template_path).namelist() if n.startswith("word/media/")]
    assert media, "no image embedded in the template"


# --- docxtpl tags present (and data rows are tags, not real names) ----------

def test_template_contains_all_jinja_tags(doc) -> None:
    text = _all_text(doc)
    for tag in (
        "{{ title }}", "{{ date }}", "{{ time }}", "{{ location }}",
        "{%tr for a in attendees %}", "{{ a.name }}", "{{ a.org }}",
        "{{ summary }}", "{%p for point in points %}", "{{ point }}",
        "{%tr for o in outcomes %}", "{{ o.task }}", "{{ o.owner }}", "{{ o.date }}",
    ):
        assert tag in text, f"missing tag: {tag}"


def test_data_rows_are_templated_not_literal_data(doc) -> None:
    # Privacy proof: attendee/outcome rows carry loop tags (marker row + var data
    # row), so no real person/org strings can be baked into the committed output.
    attendees = doc.tables[1]
    assert "{%tr for a in attendees %}" in attendees.rows[1].cells[0].text
    data_row = " ".join(c.text for c in attendees.rows[2].cells)
    assert "{{ a.name }}" in data_row and "{{ a.org }}" in data_row
    outcomes_data = " ".join(c.text for c in doc.tables[3].rows[2].cells)
    assert "{{ o.owner }}" in outcomes_data


# --- placeholder logo PNG generator ----------------------------------------

def test_build_writes_atomically_no_tmp_left(tmp_path) -> None:
    out = st.build(tmp_path / "sample.docx")
    assert out.exists()
    assert not list(tmp_path.glob("*.tmp")), "atomic-write temp file was left behind"


def test_write_solid_png_is_valid_png_of_requested_size(tmp_path) -> None:
    p = st.write_solid_png(tmp_path / "logo.png", width=120, height=40, rgb=(0x00, 0x70, 0xB9))
    data = p.read_bytes()
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    w, h = struct.unpack(">II", data[16:24])  # IHDR width/height
    assert (w, h) == (120, 40)
