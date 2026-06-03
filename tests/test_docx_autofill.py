"""Tests for auto-detect fill of an UNTOUCHED .docx (no Jinja tags).

A minimal synthetic untagged template is built in-process (no real client data),
filled, and inspected. Mirrors how the app fills a user's own uploaded template.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from meeting_minutes.docx_autofill import autofill_docx

MINUTES_MD = """# محضر اجتماع — اجتماع الاختبار (2026-06-03)

| التاريخ | الوقت | الموقع |
| --- | --- | --- |
| 2026-06-03 | 10:00 | الرياض |

## قائمة الحضور
| # | الاسم | الجهة |
| --- | --- | --- |
| 1 | فلان الأول | جهة-أ |
| 2 | فلان الثاني | جهة-ب |
| 3 | فلان الثالث | جهة-ج |

## نقاط نقاش الاجتماع
**ملخص الاجتماع**
ملخص جديد للاجتماع.

- نقطة النقاش الأولى
- نقطة النقاش الثانية

## نتائج الاجتماع
| المهام/ التوصيات | المسؤول | التاريخ المستهدف |
| --- | --- | --- |
| المهمة الأولى | فلان الأول | 2026-06-10 |
| المهمة الثانية | جهة خارجية | — |
"""


def _make_untagged_template(path) -> None:
    """A plain template (no Jinja tags) with the Arabic headers + one sample row
    per table — like a real corporate محضر template before any data is entered."""
    doc = Document()

    meta = doc.add_table(rows=3, cols=3)
    meta.cell(0, 0).merge(meta.cell(0, 2)).paragraphs[0].add_run("محضر اجتماع نموذجي")
    for cell, label in zip(meta.rows[1].cells, ("التاريخ", "الوقت", "الموقع")):
        cell.paragraphs[0].add_run(label)
    for cell, old in zip(meta.rows[2].cells, ("قديم١", "قديم٢", "قديم٣")):
        cell.paragraphs[0].add_run(old)
    doc.add_paragraph()

    attendees = doc.add_table(rows=2, cols=3)
    for cell, label in zip(attendees.rows[0].cells, ("#", "الاسم", "الجهة")):
        cell.paragraphs[0].add_run(label)
    for cell, old in zip(attendees.rows[1].cells, ("٩", "اسم عينة", "جهة عينة")):
        cell.paragraphs[0].add_run(old)
    doc.add_paragraph()

    summary = doc.add_table(rows=2, cols=1)
    summary.rows[0].cells[0].paragraphs[0].add_run("ملخص الاجتماع")
    summary.rows[1].cells[0].paragraphs[0].add_run("ملخص قديم")
    doc.add_paragraph()

    outcomes = doc.add_table(rows=2, cols=4)
    for cell, label in zip(
        outcomes.rows[0].cells, ("#", "المهام/ التوصيات", "المسؤول", "التاريخ المستهدف")
    ):
        cell.paragraphs[0].add_run(label)
    for cell, old in zip(outcomes.rows[1].cells, ("٩", "مهمة عينة", "جهة عينة", "قديم")):
        cell.paragraphs[0].add_run(old)

    doc.save(str(path))


@pytest.fixture
def filled(tmp_path):
    tpl = tmp_path / "template.docx"
    _make_untagged_template(tpl)
    return Document(BytesIO(autofill_docx(MINUTES_MD, tpl)))


def _table(doc, *needles):
    for t in doc.tables:
        header = " ".join(c.text for c in t.rows[0].cells)
        if all(n in header for n in needles):
            return t
    raise AssertionError(f"table not found: {needles}")


def _rows_text(table):
    return [[c.text.strip() for c in r.cells] for r in table.rows]


def test_meta_title_and_values_are_filled(filled) -> None:
    meta = next(
        t for t in filled.tables
        if {"التاريخ", "الوقت", "الموقع"} <= {c.text.strip() for r in t.rows for c in r.cells}
    )
    all_text = " ".join(c.text for r in meta.rows for c in r.cells)
    assert "اجتماع الاختبار" in all_text  # title updated
    assert "2026-06-03" in all_text and "10:00" in all_text and "الرياض" in all_text
    assert "قديم" not in all_text  # the template's sample values are gone


def test_attendees_rows_filled_with_index(filled) -> None:
    rows = _rows_text(_table(filled, "الاسم", "الجهة"))
    assert len(rows) == 1 + 3  # header + 3 attendees
    assert rows[1] == ["1", "فلان الأول", "جهة-أ"]
    assert [r[0] for r in rows[1:]] == ["1", "2", "3"]  # single, sequential index
    assert "عينة" not in " ".join(sum(rows, []))  # sample row removed


def test_outcomes_owner_is_resolved_org_not_person(filled) -> None:
    rows = _rows_text(_table(filled, "المسؤول"))
    assert len(rows) == 1 + 2
    assert rows[1] == ["1", "المهمة الأولى", "جهة-أ", "2026-06-10"]  # person -> org
    assert rows[2][2] == "جهة خارجية"  # non-roster owner kept verbatim
    assert "فلان الأول" not in " ".join(sum(rows, []))  # person resolved away


def test_no_doubled_index_numbers(filled) -> None:
    # The # cell must hold exactly one number (numbering was stripped on fill).
    for needles in (("الاسم", "الجهة"), ("المسؤول",)):
        for row in _rows_text(_table(filled, *needles))[1:]:
            assert row[0].isdigit() and len(row[0]) <= 2


def test_summary_paragraph_and_bullets_filled(filled) -> None:
    summary = next(t for t in filled.tables if "ملخص الاجتماع" in t.rows[0].cells[0].text)
    body = summary.rows[1].cells[0].text
    assert "ملخص جديد للاجتماع." in body
    assert "نقطة النقاش الأولى" in body and "نقطة النقاش الثانية" in body
    assert "ملخص قديم" not in body


def test_malformed_template_raises_valueerror(tmp_path) -> None:
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a real .docx zip")
    with pytest.raises(ValueError, match="not a valid .docx"):
        autofill_docx(MINUTES_MD, bad)
