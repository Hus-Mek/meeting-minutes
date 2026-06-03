"""Tests for the docx render pipeline: context build, docxtpl fill, PDF (mocked).

All data is SYNTHETIC. LibreOffice is mocked at the boundary — no subprocess, no
network, no LLM is ever exercised.
"""

from __future__ import annotations

import subprocess
from io import BytesIO

import pytest
from docx import Document

from meeting_minutes import docx_render as dr
from meeting_minutes import sample_template as st
from meeting_minutes.minutes_model import parse_minutes

SYNTHETIC_MD = """\
# محضر اجتماع — اجتماع تجريبي (2026-06-03)

| التاريخ | الوقت | الموقع |
| --- | --- | --- |
| 2026-06-03 | 10:00 | الرياض |

## قائمة الحضور
| # | الاسم | الجهة |
| --- | --- | --- |
| 1 | الحاضر الأول | جهة-أ |
| 2 | الحاضر الثاني | — |

## نقاط نقاش الاجتماع
**ملخص الاجتماع**
فقرة تمهيدية عن الاجتماع.

- نقطة النقاش الأولى
- نقطة النقاش الثانية

## نتائج الاجتماع
| المهام/ التوصيات | المسؤول | التاريخ المستهدف |
| --- | --- | --- |
| المهمة الأولى | الحاضر الأول | 2026-06-10 |
| المهمة الثانية | الحاضر الثاني | — |
| المهمة الثالثة | لجنة خارجية | — |
"""


@pytest.fixture(scope="session")
def template_path(tmp_path_factory) -> str:
    return str(st.build(tmp_path_factory.mktemp("tmpl") / "sample.docx"))


def _rendered_doc(md: str, template_path: str) -> Document:
    return Document(BytesIO(dr.render_docx(md, template_path)))


def _table_text(doc, idx: int) -> str:
    return "\n".join(c.text for r in doc.tables[idx].rows for c in r.cells)


# --- build_context ----------------------------------------------------------

def test_build_context_resolves_owner_to_org_not_person() -> None:
    ctx = dr.build_context(parse_minutes(SYNTHETIC_MD))
    owners = [o["owner"] for o in ctx["outcomes"]]
    assert owners == ["جهة-أ", "—", "لجنة خارجية"]  # org / blank→dash / verbatim non-roster
    assert all("person" not in o for o in ctx["outcomes"])


def test_build_context_empty_org_becomes_dash() -> None:
    ctx = dr.build_context(parse_minutes(SYNTHETIC_MD))
    assert ctx["attendees"][1]["org"] == "—"  # roster person with unknown الجهة


# --- fill: tags replaced, counts, owner column ------------------------------

def test_render_leaves_no_jinja_or_placeholder_tokens(template_path) -> None:
    doc = _rendered_doc(SYNTHETIC_MD, template_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for token in ("{{", "}}", "{%", "%}", "«", "»"):
        assert token not in text, f"leftover token {token!r}"


def test_render_attendee_row_count_matches_input(template_path) -> None:
    doc = _rendered_doc(SYNTHETIC_MD, template_path)
    attendees = doc.tables[1]
    assert len(attendees.rows) == 1 + 2  # header + 2 attendees


def test_render_outcome_rows_and_sequential_index(template_path) -> None:
    doc = _rendered_doc(SYNTHETIC_MD, template_path)
    outcomes = doc.tables[3]
    assert len(outcomes.rows) == 1 + 3
    index_col = [outcomes.rows[i].cells[0].text.strip() for i in range(1, 4)]
    assert index_col == ["1", "2", "3"]


def test_render_owner_column_shows_org_not_person(template_path) -> None:
    text = _table_text(_rendered_doc(SYNTHETIC_MD, template_path), 3)
    assert "جهة-أ" in text  # resolved org appears
    # the roster person's name must NOT appear in the outcomes table (it was resolved)
    assert "الحاضر الأول" not in text
    assert "لجنة خارجية" in text  # non-roster owner kept verbatim


def test_render_summary_and_each_bullet_present(template_path) -> None:
    text = _table_text(_rendered_doc(SYNTHETIC_MD, template_path), 2)
    assert "فقرة تمهيدية عن الاجتماع." in text
    assert "نقطة النقاش الأولى" in text and "نقطة النقاش الثانية" in text


# --- edge cases -------------------------------------------------------------

def test_render_with_zero_attendees_and_outcomes(template_path) -> None:
    md = "# محضر اجتماع — فارغ\n\n## قائمة الحضور\n\n## نتائج الاجتماع\n"
    doc = _rendered_doc(md, template_path)
    assert len(doc.tables[1].rows) == 1  # header only
    assert len(doc.tables[3].rows) == 1
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "{%" not in text and "{{" not in text


def test_render_escapes_special_xml_characters(template_path) -> None:
    md = (
        "# محضر اجتماع — ت\n\n"
        "## نتائج الاجتماع\n"
        "| المهام/ التوصيات | المسؤول | التاريخ المستهدف |\n"
        "| --- | --- | --- |\n"
        "| ترقية A & B < C | جهة | — |\n"
    )
    doc = _rendered_doc(md, template_path)  # must not raise / corrupt
    assert "ترقية A & B < C" in _table_text(doc, 3)


# --- PDF conversion (LibreOffice mocked) ------------------------------------

def test_docx_to_pdf_invokes_soffice(monkeypatch) -> None:
    monkeypatch.setattr(dr.shutil, "which", lambda name: "/usr/bin/soffice")
    captured = {}

    def fake_run(cmd, **kwargs):
        captured["cmd"] = cmd
        outdir = cmd[cmd.index("--outdir") + 1]
        with open(f"{outdir}/minutes.pdf", "wb") as f:
            f.write(b"%PDF-1.7 synthetic")
        return subprocess.CompletedProcess(cmd, 0, stdout=b"", stderr=b"")

    monkeypatch.setattr(dr.subprocess, "run", fake_run)
    out = dr.docx_to_pdf(b"docx-bytes")
    assert out.startswith(b"%PDF")
    assert "--headless" in captured["cmd"]
    assert captured["cmd"][captured["cmd"].index("--convert-to") + 1] == "pdf"


def test_docx_to_pdf_raises_when_soffice_missing(monkeypatch) -> None:
    monkeypatch.setattr(dr.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="not installed"):
        dr.docx_to_pdf(b"docx-bytes")


def test_docx_to_pdf_raises_when_output_missing(monkeypatch) -> None:
    monkeypatch.setattr(dr.shutil, "which", lambda name: "/usr/bin/soffice")
    monkeypatch.setattr(
        dr.subprocess, "run",
        lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, stdout=b"", stderr=b"boom"),
    )
    with pytest.raises(RuntimeError, match="conversion failed"):
        dr.docx_to_pdf(b"docx-bytes")


# --- ensure_sample_template -------------------------------------------------

def test_ensure_sample_template_generates_when_missing(monkeypatch, tmp_path) -> None:
    target = tmp_path / "nested" / "sample.docx"
    monkeypatch.setattr(dr, "SAMPLE_TEMPLATE_PATH", target)
    assert not target.exists()
    out = dr.ensure_sample_template()
    assert out == target and target.exists()
    Document(str(target))  # opens as a valid docx


def test_fill_template_rejects_non_docx_file(tmp_path) -> None:
    bad = tmp_path / "not.docx"
    bad.write_bytes(b"plain text, not a zip")
    with pytest.raises(ValueError, match="invalid or unrenderable"):
        dr.fill_template({"title": "x"}, bad)


def test_fill_template_blocks_ssti_payload(tmp_path) -> None:
    # An uploaded template attempting a Jinja SSTI/RCE chain is rejected by the
    # SandboxedEnvironment (SecurityError -> ValueError), not executed.
    from docx import Document as _Doc

    evil = _Doc()
    evil.add_paragraph("{{ ''.__class__.__mro__[1].__subclasses__() }}")
    path = tmp_path / "evil.docx"
    evil.save(str(path))
    with pytest.raises(ValueError, match="invalid or unrenderable"):
        dr.fill_template({}, path)


def test_docx_to_pdf_timeout_raises_runtimeerror(monkeypatch) -> None:
    monkeypatch.setattr(dr.shutil, "which", lambda name: "/usr/bin/soffice")

    def raise_timeout(cmd, **kwargs):
        raise subprocess.TimeoutExpired(cmd, kwargs.get("timeout", 0))

    monkeypatch.setattr(dr.subprocess, "run", raise_timeout)
    # RuntimeError (not the raw TimeoutExpired) so the endpoint maps it to 503.
    with pytest.raises(RuntimeError, match="timed out"):
        dr.docx_to_pdf(b"docx-bytes")


def test_render_docx_falls_back_to_sample_template(monkeypatch, tmp_path) -> None:
    target = tmp_path / "sample.docx"
    monkeypatch.setattr(dr, "SAMPLE_TEMPLATE_PATH", target)
    out = dr.render_docx(SYNTHETIC_MD)  # template_path=None → ensure_sample_template()
    Document(BytesIO(out))  # valid docx produced from the generated template
