"""Tests for POST /api/minutes/docx — render minutes into a .docx/PDF file.

LibreOffice is mocked for the PDF path; no LLM/network is touched. Synthetic data.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from meeting_minutes import docx_render
from meeting_minutes.sample_template import build
from meeting_minutes.web import _DOCX_MEDIA, app

MINUTES_MD = """# محضر اجتماع — اجتماع تجريبي (2026-06-03)

| التاريخ | الوقت | الموقع |
| --- | --- | --- |
| 2026-06-03 | 10:00 | الرياض |

## قائمة الحضور
| # | الاسم | الجهة |
| --- | --- | --- |
| 1 | الحاضر الأول | الجهة الأولى |
| 2 | الحاضر الثاني | الجهة الثانية |

## نقاط نقاش الاجتماع
**ملخص الاجتماع**
تم استعراض الإجراءات.

- النقطة الأولى
- النقطة الثانية

## نتائج الاجتماع
| المهام/ التوصيات | المسؤول | التاريخ المستهدف |
| --- | --- | --- |
| تعديل المصفوفة | الحاضر الأول | 2026-06-10 |
| رفع التقرير | لجنة خارجية | — |
"""


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def _outcomes_text(content: bytes) -> str:
    doc = Document(BytesIO(content))
    return "\n".join(c.text for row in doc.tables[3].rows for c in row.cells)


def test_returns_docx_file_with_no_leftover_tags(client) -> None:
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD, "filename": "minutes"})
    assert r.status_code == 200
    assert "wordprocessingml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # a .docx is a zip
    assert "attachment" in r.headers["content-disposition"]
    doc = Document(BytesIO(r.content))
    text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "{{" not in text and "{%" not in text


def test_owner_column_is_resolved_org_not_person(client) -> None:
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD})
    outcomes = _outcomes_text(r.content)
    assert "الجهة الأولى" in outcomes  # the org
    assert "الحاضر الأول" not in outcomes  # the person was resolved away
    assert "لجنة خارجية" in outcomes  # non-roster owner kept verbatim


def test_pdf_format_returns_pdf(client, monkeypatch) -> None:
    monkeypatch.setattr(docx_render, "docx_to_pdf", lambda data: b"%PDF-1.7 stub")
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD, "format": "pdf"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")


def test_pdf_unavailable_returns_503(client, monkeypatch) -> None:
    def boom(data):
        raise RuntimeError("LibreOffice (soffice) is not installed; cannot convert to PDF")

    monkeypatch.setattr(docx_render, "docx_to_pdf", boom)
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD, "format": "pdf"})
    assert r.status_code == 503
    assert "LibreOffice" in r.json()["error"]


def test_invalid_format_returns_400(client) -> None:
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD, "format": "rtf"})
    assert r.status_code == 400


def test_uploaded_template_is_used(client, tmp_path) -> None:
    tpl = build(tmp_path / "sample.docx")
    files = {"template": ("mytemplate.docx", tpl.read_bytes(), _DOCX_MEDIA)}
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD}, files=files)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    assert "الجهة الأولى" in _outcomes_text(r.content)


def test_malformed_uploaded_template_returns_400(client) -> None:
    files = {"template": ("broken.docx", b"this is not a .docx zip", _DOCX_MEDIA)}
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD}, files=files)
    assert r.status_code == 400
    assert "template" in r.json()["error"]


def test_minutes_is_required(client) -> None:
    r = client.post("/api/minutes/docx", data={})
    assert r.status_code == 422  # FastAPI request validation


def test_content_disposition_carries_utf8_arabic_name(client) -> None:
    r = client.post("/api/minutes/docx", data={"minutes": MINUTES_MD, "filename": "محضر"})
    assert r.status_code == 200
    cd = r.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd and cd.rstrip().endswith(".docx")
